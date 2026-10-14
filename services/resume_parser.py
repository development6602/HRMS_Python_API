import io
import re
from datetime import datetime
from fastapi import HTTPException, status
from pypdf import PdfReader
from docx import Document
from services.extractor_service import extract_text_via_ocr
import pymupdf

MONTH_MAP = {
    'jan': 1, 'january': 1, 'feb': 2, 'february': 2, 'mar': 3, 'march': 3,
    'apr': 4, 'april': 4, 'may': 5, 'jun': 6, 'june': 6, 'jul': 7, 'july': 7,
    'aug': 8, 'august': 8, 'sep': 9, 'sept': 9, 'september': 9, 'oct': 10,
    'october': 10, 'nov': 11, 'november': 11, 'dec': 12, 'december': 12
}

def extract_text_from_pdf(content: bytes) -> str:
    """PDF se text extract karta hai (PyMuPDF with OCR Fallback)."""
    try:
        doc = pymupdf.open(stream=content, filetype="pdf")
        text = "\n".join([page.get_text() for page in doc]).strip()
    except Exception:
        pdf_reader = PdfReader(io.BytesIO(content))
        text_pieces = [page.extract_text() or "" for page in pdf_reader.pages]
        text = "\n".join([t for t in text_pieces if t]).strip()
    
    # Agar digital text blank mila, toh Local OCR se nikaalo
    if not text:
        text = extract_text_via_ocr(content)
        
    return text

def extract_text_from_docx(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = list(dict.fromkeys([c.text.strip() for c in row.cells if c.text.strip()]))
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)

def extract_text_from_binary_doc(content: bytes) -> str:
    if content.startswith(b'PK\x03\x04'):
        return extract_text_from_docx(content)
    utf16 = [m.decode('utf-16le').strip() for m in re.findall(b'(?:[\x20-\x7E]\x00){3,}', content)]
    return "\n".join([t for t in utf16 if t])

def extract_resume_text(content: bytes, filename: str) -> str:
    lower_name = (filename or "").lower()
    if lower_name.endswith(".pdf") or content.startswith(b'%PDF'):
        return extract_text_from_pdf(content)
    elif lower_name.endswith(".docx") or content.startswith(b'PK\x03\x04'):
        return extract_text_from_docx(content)
    else:
        return extract_text_from_binary_doc(content)

def parse_date_to_month_index(token: str) -> int:
    now = datetime.now()
    clean = token.strip().lower()
    if any(p in clean for p in ['present', 'current', 'continuing', 'continue']):
        return now.year * 12 + now.month
    slash = re.search(r'\b(0?[1-9]|1[0-2])[\/\-](\d{4})\b', clean)
    if slash:
        return int(slash.group(2)) * 12 + int(slash.group(1))
    multi = re.search(r'([a-z]+)[-\/]([a-z]+)\s*(\d{4})', clean)
    if multi and multi.group(2) in MONTH_MAP:
        return int(multi.group(3)) * 12 + MONTH_MAP[multi.group(2)]
    word = re.search(r'([a-z]+)\s*(\d{4})', clean)
    if word and word.group(1) in MONTH_MAP:
        return int(word.group(2)) * 12 + MONTH_MAP[word.group(1)]
    year = re.search(r'\b(20\d{2}|19\d{2})\b', clean)
    if year:
        return int(year.group(1)) * 12 + 6
    return 0

def parse_candidate_email(text: str) -> str:
    m = re.search(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', text)
    return m.group(0).rstrip('.,;:') if m else ""

def parse_candidate_phone(text: str) -> str:
    matches = re.findall(r'(?:\+?\d{1,3}[\s\-]?)?(?:\(?\d{2,5}\)?[\s\-]?)?\d{4,5}[\s\-]?\d{4,5}', text)
    for m in matches:
        digits = re.sub(r'\D', '', m)
        if 10 <= len(digits) <= 12:
            return m.strip()
    return ""

def parse_candidate_name(text: str) -> tuple[str, str]:
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    blacklist = {
        'resume', 'cv', 'curriculum', 'contact', 'summary', 'profile', 'education',
        'experience', 'projects', 'skills', 'portfolio', 'activities', 'achievements',
        'certifications', 'languages', 'declaration', 'ai', 'ml', 'software', 'solutions',
        'technologies', 'technology', 'infotech', 'services', 'systems', 'consulting',
        'pvt', 'ltd', 'corp', 'inc', 'llc', 'developer', 'engineer', 'intern',
        'undergraduate', 'student', 'graduate', 'linkedin', 'linkdin', 'github', 'leetcode'
    }
    for line in lines[:15]:
        line_clean = re.sub(r'\(.*?\)', '', line).strip()
        clean = re.sub(r'[^a-zA-Z\s]', '', line_clean).strip()
        words = clean.split()
        if 2 <= len(words) <= 4 and not any(w.lower() in blacklist for w in words):
            if '@' not in line and not re.search(r'\+\d{2,}', line) and not any(ch.isdigit() for ch in line):
                return words[0].title(), ' '.join(words[1:]).title()
    for line in lines[:5]:
        line_clean = re.sub(r'\(.*?\)', '', line).strip()
        clean = re.sub(r'[^a-zA-Z\s]', '', line_clean).strip()
        words = clean.split()
        if len(words) == 1 and words[0].lower() not in blacklist:
            return words[0].title(), ''
    return '', ''

def normalize_spaced_letters(text: str) -> str:
    chunks = re.split(r'\s{2,}', text.strip())
    fixed = []
    for c in chunks:
        toks = c.split()
        if len(toks) >= 3 and all(len(t) == 1 for t in toks):
            fixed.append(''.join(toks))
        else:
            fixed.append(c)
    return " ".join(fixed)

def parse_candidate_position(text: str) -> str:
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    role_keywords = ['developer', 'engineer', 'programmer', 'analyst', 'scientist', 'executive', 'lead', 'intern', 'specialist', 'consultant', 'architect', 'manager']

    # 1. Header lines (Lines 1 to 10) me direct designation check
    for line in lines[1:10]:
        norm_line = normalize_spaced_letters(line)
        lower = norm_line.lower()
        if '@' not in line and not re.search(r'\+\d{2,}', line) and 'http' not in lower and 'github' not in lower and 'linkedin' not in lower:
            parts = [p.strip() for p in norm_line.split('|')]
            for p in parts:
                clean_p = re.sub(r'[^a-zA-Z\/\-\s\&]', '', p).strip()
                lower_p = clean_p.lower()
                if any(k in lower_p for k in role_keywords):
                    if not any(exp in lower_p for exp in ['experience', 'years', 'yrs', 'skills', 'summary', 'languages', 'gujarat', 'surat', 'india']):
                        if 3 <= len(clean_p) <= 40:
                            return clean_p.title()

    # 2. EXPERIENCE section ke andar se designation nikalein
    exp_match = re.search(
        r'(?:WORK\s+EXPERIENCE|PROFESSIONAL\s+EXPERIENCE|EMPLOYMENT\s+HISTORY|INTERNSHIPS?|EXPERIENCE)\s*[:\n\r]+'
        r'([\s\S]+?)'
        r'(?=\n\s*(?:PROJECTS?|PROJECT\s+DETAILS|ACADEMIC\s+QUALIFICATIONS|EDUCATION|TECHNICAL\s+SKILLS|SKILLS|CERTIFICATIONS?|CERTIFICATION\s+COURSES?|ACHIEVEMENTS?|AWARDS?|PUBLICATIONS?|LANGUAGES|DECLARATION)\s*[:\n\r]|$)',
        text,
        re.I
    )
    if exp_match:
        exp_lines = [l.strip() for l in exp_match.group(1).splitlines() if l.strip()]
        for l in exp_lines:
            if l.startswith(('•', '-', '*', '—')):
                continue
            cleaned_l = re.sub(r'\(?\d{1,2}[\/\-]\d{4}.*$', '', l).strip()
            cleaned_l = re.sub(r'\(?\b[a-zA-Z]+\s+\d{4}.*$', '', cleaned_l).strip()
            lower_l = cleaned_l.lower()
            if any(k in lower_l for k in role_keywords):
                clean_title = re.split(r'\s{3,}|\t|\||,|\b(?:India|Remote|Surat|Delhi|Vadodara|Ahmedabad|Mumbai|Pune|Bangalore)\b', cleaned_l)[0].strip()
                clean_title = re.sub(r'[^a-zA-Z\/\-\s\&]', '', clean_title).strip()
                if 3 <= len(clean_title) <= 40:
                    return clean_title

    # 3. Standard technical roles fallback
    common = [
        "Full Stack Developer", "Fullstack Developer", "Full Stack Engineer", "Fullstack Engineer",
        "MERN Stack Developer", "AI/ML Developer", "AI-ML Developer", "AI/ML Engineer", "AI-ML Engineer",
        "Machine Learning Developer", "Machine Learning Engineer", "Data Scientist", "Python Developer",
        "Python Engineer", "Frontend Developer", "Backend Developer", "Software Engineer", "Software Developer",
        "Web Developer", "Data Analyst", "DevOps Engineer", "React Developer", "BDE"
    ]
    for r in common:
        if re.search(r'(?<!\w)' + r'\s+'.join(re.escape(w) for w in r.split()) + r'(?!\w)', text[:2000], re.I):
            return r

    return ""

def parse_candidate_experience(text: str) -> str:
    """
    Tier 1: WORK EXPERIENCE / INTERNSHIPS section se exact date duration calculate karta hai.
            (Interval Merging ke saath taaki parallel projects double-count na hon).
    Tier 2 (Fallback): Text-based "X years experience" ya "Fresher" check karta hai.
    """
    # ========================================================
    # TIER 1: EXPERIENCE SECTION DATES CALCULATION
    # ========================================================
    # Fix 1: Chipke huye month words ko alag karein (e.g. TechnologiesJan -> Technologies Jan)
    normalized_text = re.sub(
        r'([a-zA-Z])(January|February|March|April|May|June|July|August|September|October|November|December|Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\b',
        r'\1 \2',
        text,
        flags=re.IGNORECASE
    )

    # Fix 2: Lookahead ko sirf nayi line (\n\s*) wale section headers par match karein
    # (ACHIEVEMENTS?, CERTIFICATIONS?, CERTIFICATION COURSES? sabhi singular/plural cover hain)
    exp_section_match = re.search(
        r'(?:WORK\s+EXPERIENCE|PROFESSIONAL\s+EXPERIENCE|EMPLOYMENT\s+HISTORY|INTERNSHIPS?|EXPERIENCE)\s*[:\n\r]+'
        r'([\s\S]+?)'
        r'(?=\n\s*(?:PROJECTS?|PROJECT\s+DETAILS|ACADEMIC\s+QUALIFICATIONS|EDUCATION|TECHNICAL\s+SKILLS|SKILLS|CERTIFICATIONS?|CERTIFICATION\s+COURSES?|ACHIEVEMENTS?|AWARDS?|PUBLICATIONS?|LANGUAGES|DECLARATION)\s*[:\n\r]|$)',
        normalized_text,
        re.IGNORECASE
    )

    token_rx = r'(?:(?:\d{1,2}[\/\-])?(?:[a-zA-Z]+(?:-[a-zA-Z]+)?\s*)?\d{4}|present|current|continuing|continue)'
    date_range_pattern = rf'({token_rx})\s*[-—–\ufffd"\'to~]+\s*({token_rx})'

    ranges = []
    if exp_section_match:
        exp_text = exp_section_match.group(1)
        ranges = re.findall(date_range_pattern, exp_text, re.IGNORECASE)

    # Dynamic Fallback: Agar Canva/Figma floating box ki wajah se header ke baad dates na milein,
    # toh work/intern context wali dates nikaalein (aur Education context ko skip karein)
    if not ranges:
        for m in re.finditer(date_range_pattern, normalized_text, re.IGNORECASE):
            s_str, e_str = m.group(1), m.group(2)
            ctx = normalized_text[max(0, m.start() - 150):min(len(normalized_text), m.end() + 150)].lower()
            if any(edu in ctx for edu in ['school', 'college', 'university', 'bachelor', 'b.e.', 'b.tech', 'diploma', 'board', 'cgpa', 'percentage']):
                continue
            if any(w in ctx for w in ['intern', 'developer', 'engineer', 'work', 'experience', 'technologies', 'solutions', 'full stack', 'backend', 'ai/ml', 'analyst', 'systems']):
                ranges.append((s_str, e_str))

    intervals = []
    for start_str, end_str in ranges:
        start_idx = parse_date_to_month_index(start_str)
        end_idx = parse_date_to_month_index(end_str)
        if start_idx > 0 and end_idx >= start_idx and (end_idx - start_idx) <= 180:
            intervals.append((start_idx, end_idx + 1))

    # ✨ Overlapping Date Intervals ko Merge karein (Loop ke bahar 4 spaces)
    if intervals:
        intervals.sort(key=lambda x: x[0])
        merged = []
        for s, e in intervals:
            if not merged or s > merged[-1][1]:
                merged.append([s, e])
            else:
                merged[-1][1] = max(merged[-1][1], e)

        # Merged continuous working months ka sum
        total_months = sum(e - s for s, e in merged)

        if total_months >= 12:
            years = round(total_months / 12, 1)
            return f"{int(years) if years.is_integer() else years} Years"
        elif total_months > 0:
            return f"{total_months} Months"

    # ========================================================
    # TIER 2: FALLBACK (AGAR DATES NA MILEIN TO PURANA LOGIC)
    # ========================================================
    lower_text = text.lower()

    # 1. Direct explicit tenure mention (e.g. "2.5 years of experience")
    tenure_match = re.search(
        r'(\d+(?:\.\d+)?)\+?\s*(?:years?|yrs?)(?:\s*(?:of)?\s*experience)?',
        text,
        re.IGNORECASE
    )
    if tenure_match:
        return f"{tenure_match.group(1)} Years"

    # 2. Months mention (e.g. "6 months experience")
    month_match = re.search(
        r'(\d+)\+?\s*(?:months?|mos?)(?:\s*(?:of)?\s*experience)?',
        text,
        re.IGNORECASE
    )
    if month_match:
        m_val = int(month_match.group(1))
        return f"{round(m_val / 12, 1)} Years" if m_val >= 12 else f"{m_val} Months"

    # 3. Direct Fresher keywords
    if "fresher" in lower_text or "entry level" in lower_text or "recent graduate" in lower_text:
        return "Fresher"

    # 4. College pursuing student check
    year_matches = [int(y) for y in re.findall(r'\b(20[12]\d)\b', text)]
    if year_matches and max(year_matches) >= datetime.now().year:
        return "Fresher"

    return "Fresher"


def parse_uploaded_resume_file(content: bytes, filename: str) -> dict:
    lower_name = (filename or "").lower()
    if not any(lower_name.endswith(ext) for ext in (".pdf", ".docx", ".doc")):
        raise HTTPException(status_code=400, detail=f"Unsupported format '{filename}'. Allowed: .pdf, .docx, .doc")
    if not content:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")
    resume_text = extract_resume_text(content, filename).strip()
    if not resume_text:
        raise HTTPException(status_code=422, detail="Could not extract readable text from uploaded file.")
    first_name, last_name = parse_candidate_name(resume_text)
    return {
        "success": True,
        "data": {
            "firstName": first_name,
            "lastName": last_name,
            "email": parse_candidate_email(resume_text),
            "phone": parse_candidate_phone(resume_text),
            "position": parse_candidate_position(resume_text),
            "experience": parse_candidate_experience(resume_text)
        },
        "message": "Resume parsed successfully"
    }
