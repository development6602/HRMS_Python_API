import io
import os
import requests
from fastapi import HTTPException
from pypdf import PdfReader
from docx import Document
from services.candidate_service import fetch_candidate_from_nodejs
from urllib.parse import urlparse
from rapidocr_onnxruntime import RapidOCR
import pymupdf

_ocr_engine = None

def get_ocr_engine():
    global _ocr_engine
    if _ocr_engine is None:
        _ocr_engine = RapidOCR()
    return _ocr_engine

def extract_text_via_ocr(content: bytes) -> str:
    """Scanned / Image-only PDF se local OCR dwara text extract karta hai."""
    try:
        doc = pymupdf.open(stream=content, filetype="pdf")
        engine = get_ocr_engine()
        lines = []
        for page in doc:
            pix = page.get_pixmap(dpi=150)
            img_bytes = pix.tobytes("png")
            result, _ = engine(img_bytes)
            if result:
                for item in result:
                    lines.append(item[1])
        return "\n".join(lines).strip()
    except Exception as e:
        return ""

def extract_text_from_pdf_bytes(content: bytes) -> dict:
    """PDF file bytes se text aur page count extract karta hai (PyMuPDF with OCR Fallback)."""
    try:
        doc = pymupdf.open(stream=content, filetype="pdf")
        pages_text = [page.get_text() for page in doc]
        cleaned_text = "\n".join(pages_text).strip()
        total_pages = len(doc)
    except Exception:
        pdf_file = io.BytesIO(content)
        reader = PdfReader(pdf_file)
        full_text = ""
        for page in reader.pages:
            full_text += (page.extract_text() or "") + "\n"
        cleaned_text = full_text.strip()
        total_pages = len(reader.pages)
    
    # Agar digital text 0 nikla (Scanned/Photo PDF), toh Local OCR chalao
    if not cleaned_text:
        cleaned_text = extract_text_via_ocr(content)

    return {
        "fileType": "pdf",
        "totalPages": total_pages,
        "fullText": cleaned_text,
        "characterCount": len(cleaned_text)
    }

def extract_text_from_docx_bytes(content: bytes) -> dict:
    """DOCX file bytes ke paragraphs aur tables se text extract karta hai."""
    doc = Document(io.BytesIO(content))
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)
            
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            unique_cells = list(dict.fromkeys(cells))
            if unique_cells:
                lines.append(" | ".join(unique_cells))
                
    cleaned_text = "\n".join(lines).strip()
    return {
        "fileType": "docx",
        "totalPages": None,
        "fullText": cleaned_text,
        "characterCount": len(cleaned_text)
    }


# def extract_text_from_resume_url(resume_url: str, file_name: str = "") -> dict:
#     """Smart Extractor: PDF aur DOCX dono ko detect karta hai."""
#     try:
#         download_response = requests.get(resume_url, timeout=15)
#         if download_response.status_code != 200:
#             raise HTTPException(
#                 status_code=400,
#                 detail=f"Could not download resume from URL (Status: {download_response.status_code})"
#             )
            
#         content = download_response.content
#         lower_name = (file_name or "").lower()
def extract_text_from_resume_url(resume_url: str, file_name: str = "") -> dict:
    """
    Smart Extractor: PDF aur DOCX dono ko automatically detect karke text nikalta hai.
    - 'website_form' (Cloudinary https://): Direct download karega (As-it-is).
    - 'admin_panel' (Local /api/files/...): Server Origin + JWT Token add karke Node.js se download karega.
    """
    try:
        download_url = resume_url
        headers = {}

        # 🌐 Case 2: Agar Admin Panel se upload hua local resume hai (Relative path starts with '/')
        if resume_url.startswith("/"):
            hrms_base = os.getenv("HRMS_BASE_URL", "http://localhost:5000/api")
            hrms_token = os.getenv("HRMS_AUTH_TOKEN", "")
            
            # Base URL se origin domain extract karein (e.g. http://localhost:5000 ya http://192.168.1.206)
            parsed_base = urlparse(hrms_base)
            server_origin = f"{parsed_base.scheme}://{parsed_base.netloc}"
            download_url = f"{server_origin}{resume_url}"
            
            # Node.js backend ko access karne ke liye JWT Auth Token add karein
            if hrms_token:
                headers["Authorization"] = f"Bearer {hrms_token}"

        # 🚀 Resume download karein (Cloudinary ke liye direct, Admin panel ke liye with token)
        download_response = requests.get(download_url, headers=headers, timeout=15)
        if download_response.status_code != 200:
            raise HTTPException(
                status_code=400,
                detail=f"Could not download resume from URL (Status: {download_response.status_code})"
            )
            
        content = download_response.content
        lower_name = (file_name or "").lower()

        
        if content.startswith(b'%PDF') or lower_name.endswith('.pdf'):
            return extract_text_from_pdf_bytes(content)
        elif content.startswith(b'PK\x03\x04') or lower_name.endswith('.docx') or lower_name.endswith('.doc'):
            return extract_text_from_docx_bytes(content)
        else:
            try:
                return extract_text_from_pdf_bytes(content)
            except Exception:
                return extract_text_from_docx_bytes(content)
                
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error extracting text from resume: {str(e)}")


def get_candidate_and_extract_resume_data(application_id: str) -> dict:
    """Candidate details + Resume text ka structured response banata hai."""
    raw_response = fetch_candidate_from_nodejs(application_id)
    candidate_data = raw_response.get("data") or {}
    
    resume_info = candidate_data.get("resume") or {}
    resume_url = resume_info.get("resumeDisplayUrl") or resume_info.get("fileUrl")
    file_name = resume_info.get("fileName") or ""
    
    if not resume_url:
        raise HTTPException(status_code=404, detail="No resume URL found for this candidate")
        
    extracted_resume_data = extract_text_from_resume_url(resume_url, file_name)
    
    return {
        "success": True,
        "applicationId": candidate_data.get("applicationId"),
        "name": candidate_data.get("name"),
        "email": candidate_data.get("email"),
        "phone": candidate_data.get("phone"),
        "position": candidate_data.get("position"),
        "experience": candidate_data.get("experience"),
        "status": candidate_data.get("status"),
        "resume": {
            "fileName": file_name,
            "fileUrl": resume_url,
            "fileType": extracted_resume_data.get("fileType"),
            "totalPages": extracted_resume_data.get("totalPages"),
            "characterCount": extracted_resume_data.get("characterCount"),
            "rawExtractedText": extracted_resume_data.get("fullText")
        },
        "message": "Candidate resume text extracted successfully"
    }
